"""Таблица разметки v1 в xlsx: перевод на русский, список ok/fix/drop, прогресс.

    pip install -e ".[review]"
    python scripts/build_review_xlsx.py data/review_v1_cos.csv \
        data/review_v1_translation_ru.csv data/review_v1_cos.xlsx
"""
import csv, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

src, tr_path, dst = sys.argv[1], sys.argv[2], sys.argv[3]
rows = list(csv.DictReader(open(src, encoding="utf-8")))
tr = {r["id"]: r for r in csv.DictReader(open(tr_path, encoding="utf-8"))}

# Метка переводится кодом, а не моделью: это то, что проверяется.
WORDS = {"black": "чёрный", "white": "белый", "blue": "синий", "green": "зелёный",
         "grey": "серый", "pink": "розовый", "red": "красный", "silver": "серебристый",
         "beginner": "начальный", "intermediate": "средний", "pro": "профи",
         "bag": "сумка", "balls": "мячи", "racket": "ракетка", "shoes": "кроссовки"}
KEYS = {"family": "семейство", "color": "цвет", "size_label": "размер",
        "category": "категория", "level": "уровень", "pack_size": "в упаковке"}


def label_ru(s: str) -> str:
    s = re.sub(r"\b(" + "|".join(WORDS) + r")\b", lambda m: WORDS[m.group(1)], s)
    s = re.sub(r"\b(" + "|".join(KEYS) + r")\b", lambda m: KEYS[m.group(1)], s)
    for a, b in [("ships in ≤ ", "доставка ≤ "), (" days", " дн."), ("final sale", "без возврата"),
                 ("returnable", "с возвратом"), ("total ≤ ", "итого ≤ "),
                 ("should ask first", "сначала спросить человека"),
                 ("buy nothing", "ничего не покупать"), ("price ≤ ", "цена ≤ "),
                 (" ∈ ", " — любой из "), (" × [", " шт. × [")]:
        s = s.replace(a, b)
    s = re.sub(r"\bnot ", "не ", s)
    return s


F = "Arial"
bold = Font(name=F, bold=True)
norm = Font(name=F)
grey = Font(name=F, color="808080")
wrap = Alignment(wrap_text=True, vertical="top")
head_fill = PatternFill("solid", fgColor="D9D9D9")
input_fill = PatternFill("solid", fgColor="FFFF99")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# (заголовок, ширина, откуда значение)
cols = [
    ("id", 24, lambda r: r["id"]),
    ("шаблон", 16, lambda r: r["template"]),
    ("язык", 6, lambda r: r["lang"]),
    ("просьба (как её видит агент)", 45, lambda r: r["request"]),
    ("просьба по-русски", 45, lambda r: r["request"] if r["lang"] == "ru" else tr[r["id"]]["request_ru"]),
    ("что человек имеет в виду", 45, lambda r: tr[r["id"]]["persona_ru"]),
    ("правильная покупка (метка)", 40, lambda r: label_ru(r["label"])),
    ("правильная покупка (артикул)", 30, lambda r: r["right_purchase"]),
    ("ловушка", 36, lambda r: tr[r["id"]]["trap_ru"]),
    ("verdict", 10, lambda r: r["verdict"]),
    ("comment", 40, lambda r: r["comment"]),
    ("split", 8, lambda r: r["split"]),
    ("persona (оригинал)", 45, lambda r: r["persona"]),
    ("label (оригинал)", 40, lambda r: r["label"]),
    ("trap (оригинал)", 36, lambda r: r["trap"]),
]
V = 10  # номер колонки verdict
VL = "J"
n = len(rows) + 1

wb = Workbook()
ins = wb.active
ins.title = "Инструкция"
lines = [
    ("Разметка датасета AsAsked v1", None),
    ("", None),
    ("Что делать", None),
    ("Для каждой строки на листе «Задачи» проверьте по смыслу:", None),
    ("правильная покупка действительно правильная для этой просьбы, а ловушка — действительно ошибка?", None),
    ("Естественность английского текста оценивать не нужно — её проверяет модель отдельно.", None),
    ("", None),
    ("Заполнять только жёлтые колонки: verdict (выбор из списка) и comment.", None),
    ("ok — всё верно", None),
    ("fix — задача годится, но надо поправить; в comment — что именно", None),
    ("drop — выбросить; в comment — почему", None),
    ("Остальные колонки не менять. Колонку id не трогать — по ней правки вернутся в датасет.", None),
    ("", None),
    ("Откуда перевод", None),
    ("Просьба, «что имеет в виду» и ловушка переведены моделью Claude Sonnet 5 — перевод может ошибаться.", None),
    ("Метка переведена кодом, дословно. Если перевод и оригинал расходятся, верен оригинал (серые колонки справа).", None),
    ("", None),
    ("Пример заполнения", None),
    ("verdict = fix, comment = «он просит две банки, в метке должно быть ×2»", None),
    ("verdict = drop, comment = «просьба двусмысленная, правильного ответа нет»", None),
    ("", None),
    ("Прогресс", None),
    ("Всего задач", f"=COUNTA(Задачи!A2:A{n})"),
    ("Размечено", f"=COUNTA(Задачи!{VL}2:{VL}{n})"),
    ("ok", f'=COUNTIF(Задачи!{VL}2:{VL}{n},"ok")'),
    ("fix", f'=COUNTIF(Задачи!{VL}2:{VL}{n},"fix")'),
    ("drop", f'=COUNTIF(Задачи!{VL}2:{VL}{n},"drop")'),
]
titles = {"Разметка датасета AsAsked v1", "Что делать", "Откуда перевод", "Пример заполнения", "Прогресс"}
for i, (a, b) in enumerate(lines, 1):
    ins.cell(i, 1, a).font = bold if a in titles else norm
    if b is not None:
        ins.cell(i, 2, b).font = norm
total_row = next(i for i, (a, _) in enumerate(lines, 1) if a == "Всего задач")
ins.cell(len(lines) + 1, 1, "Осталось").font = norm
ins.cell(len(lines) + 1, 2, f"=B{total_row}-B{total_row + 1}").font = norm
ins["A1"].font = Font(name=F, bold=True, size=14)
ins.column_dimensions["A"].width = 100
ins.column_dimensions["B"].width = 12

ws = wb.create_sheet("Задачи")
for c, (name, w, _) in enumerate(cols, 1):
    cell = ws.cell(1, c, name)
    cell.font, cell.border, cell.alignment = bold, box, wrap
    cell.fill = input_fill if name in ("verdict", "comment") else head_fill
    ws.column_dimensions[cell.column_letter].width = w
for r, row in enumerate(rows, 2):
    for c, (name, _, get) in enumerate(cols, 1):
        cell = ws.cell(r, c, get(row))
        cell.alignment, cell.border = wrap, box
        cell.font = grey if c > 11 else norm
        if name in ("verdict", "comment"):
            cell.fill = input_fill
assert ws.cell(1, V).value == "verdict"
dv = DataValidation(type="list", formula1='"ok,fix,drop"', allow_blank=True,
                    showErrorMessage=True, errorTitle="verdict", error="Только ok, fix или drop")
ws.add_data_validation(dv)
dv.add(f"{VL}2:{VL}{n}")
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{ws.cell(1, len(cols)).column_letter}{n}"
wb.save(dst)
print(len(rows), "rows ->", dst)
